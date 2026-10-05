import re
import docx

md_path = r"C:\Users\Goutham\.gemini\antigravity-ide\brain\494f7a94-1825-42e2-95a6-de48a89abf9a\literature_survey_ieee_scopus.md"
docx_path = r"d:\AI_Text_Checker\crossref.docx"

with open(md_path, "r", encoding="utf-8") as f:
    content = f.read()

# Split by Paper headings
paper_blocks = re.split(r'## 📄 Paper \d+:', content)[1:]

papers = []
for i, block in enumerate(paper_blocks, 1):
    # Extract Title
    title_match = re.search(r'-\s*\*\*Title\*\*:\s*(.+)', block)
    title = title_match.group(1).strip() if title_match else f"Paper {i}"
    
    # Extract Abstract
    abs_match = re.search(r'### Abstract\s*\n+>\s*(.+?)(?=\n\n|\n###|\Z)', block, re.DOTALL)
    if abs_match:
        abstract = abs_match.group(1).replace("\n> ", " ").replace("\n>", " ").strip()
    else:
        abstract = ""
        
    # Extract IEEE Citation Text
    cit_match = re.search(r'### IEEE Citation Text\s*\n+```text\s*\n(.+?)\n```', block, re.DOTALL)
    if cit_match:
        citation = cit_match.group(1).strip()
    else:
        citation = ""
        
    papers.append({
        "num": i,
        "title": title,
        "abstract": abstract,
        "citation": citation
    })

print(f"Extracted {len(papers)} papers.")
for p in papers[:3]:
    print(f"--- Paper {p['num']} ---")
    print("Title:", p["title"])
    print("Abstract:", p["abstract"][:100] + "...")
    print("Citation:", p["citation"])

# Create clean Word doc
doc = docx.Document()

# Set title/heading for document
doc.add_heading("IEEE Research Paper Literature Survey & Citation Library", level=0)

for p in papers:
    # 1. Title of research paper
    p_title = doc.add_paragraph()
    run_num = p_title.add_run(f"Paper {p['num']}: ")
    run_num.bold = True
    run_title = p_title.add_run(p["title"])
    run_title.bold = True
    
    # 2. Abstract
    p_abs_lbl = doc.add_paragraph()
    run_abs_lbl = p_abs_lbl.add_run("Abstract: ")
    run_abs_lbl.bold = True
    p_abs_lbl.add_run(p["abstract"])
    
    # 3. Citation
    p_cit_lbl = doc.add_paragraph()
    run_cit_lbl = p_cit_lbl.add_run("Citation: ")
    run_cit_lbl.bold = True
    p_cit_lbl.add_run(p["citation"])
    
    # Empty paragraph separator
    doc.add_paragraph()

doc.save(docx_path)
print(f"Successfully saved {len(papers)} papers to {docx_path}")
