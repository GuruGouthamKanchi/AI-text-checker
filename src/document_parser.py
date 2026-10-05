import os
import sys
import re
import math
import zipfile
import subprocess
import shutil

class DocumentParsingError(Exception):
    """Custom exception raised for document parsing and formatting errors."""
    pass

class PasswordProtectedError(DocumentParsingError):
    """Exception raised when a file is password-protected/encrypted."""
    pass

class MissingDependencyError(DocumentParsingError):
    """Exception raised when a required system dependency (e.g. LibreOffice) is missing."""
    pass

def detect_file_type(file_path: str) -> str:
    """
    Detects the file type by reading the first 8 bytes of the file (magic bytes).
    Falls back to checking the file extension if magic bytes are ambiguous.
    """
    try:
        with open(file_path, 'rb') as f:
            header = f.read(8)
    except Exception as e:
        raise DocumentParsingError(f"Failed to read file header: {str(e)}")

    # PDF magic bytes: %PDF-
    if header.startswith(b'%PDF-'):
        return 'pdf'

    # OLE Compound Document (legacy DOC): D0 CF 11 E0 A1 B1 1A E1
    if header.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        return 'doc'

    # Zip magic bytes: PK\x03\x04 (DOCX is a zip containing [Content_Types].xml and word/)
    if header.startswith(b'PK\x03\x04'):
        if zipfile.is_zipfile(file_path):
            try:
                with zipfile.ZipFile(file_path) as z:
                    if any('word/document.xml' in name for name in z.namelist()):
                        return 'docx'
            except Exception:
                pass

    # Fallback to extension check if magic bytes are inconclusive (e.g. converted or stripped files)
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return 'pdf'
    elif ext == '.docx':
        return 'docx'
    elif ext == '.doc':
        return 'doc'

    raise DocumentParsingError(
        "Unsupported file format. Only academic manuscripts in PDF (.pdf), Word (.docx), or legacy Word (.doc) formats are supported."
    )

import base64

def extract_pdf_visual_and_text(file_path: str) -> dict:
    """
    Extracts text, page rendering images (base64 PNG), and embedded figures/tables
    from an academic PDF using PyMuPDF (fitz).
    """
    try:
        import fitz
        doc = fitz.open(file_path)
        if doc.is_encrypted:
            raise PasswordProtectedError("The PDF document is password-protected. Please remove the password and try again.")
            
        page_count = len(doc)
        text_pages = []
        page_images = []
        extracted_figures = []
        
        for i, page in enumerate(doc):
            # 1. Extract text
            page_text = page.get_text("text")
            if page_text:
                text_pages.append(page_text)
                
            # 2. Render high-res visual page image (base64 PNG)
            pix = page.get_pixmap(dpi=150)
            img_bytes = pix.tobytes("png")
            base64_img = f"data:image/png;base64,{base64.b64encode(img_bytes).decode('utf-8')}"
            page_images.append({
                "page_number": i + 1,
                "image_data": base64_img,
                "width": pix.width,
                "height": pix.height
            })
            
            # 3. Extract embedded diagrams, charts, figures & tables
            try:
                image_list = page.get_images(full=True)
                for img_index, img_info in enumerate(image_list[:4]):
                    xref = img_info[0]
                    base_image = doc.extract_image(xref)
                    image_bytes = base_image["image"]
                    image_ext = base_image["ext"]
                    if base_image["width"] >= 60 and base_image["height"] >= 60:
                        fig_b64 = f"data:image/{image_ext};base64,{base64.b64encode(image_bytes).decode('utf-8')}"
                        extracted_figures.append({
                            "id": f"fig-{i+1}-{img_index+1}",
                            "page_number": i + 1,
                            "image_data": fig_b64,
                            "width": base_image["width"],
                            "height": base_image["height"]
                        })
            except Exception:
                pass
                
        full_text = "\n\n".join(text_pages)
        if not full_text.strip():
            raise DocumentParsingError("No readable text could be extracted from PDF.")
            
        return {
            "text": full_text,
            "page_count": page_count,
            "page_images": page_images,
            "extracted_figures": extracted_figures
        }
    except PasswordProtectedError:
        raise
    except Exception as e:
        text, p_count = extract_text_from_pdf(file_path)
        return {
            "text": text,
            "page_count": p_count,
            "page_images": [],
            "extracted_figures": []
        }

def extract_text_from_pdf(file_path: str) -> tuple[str, int]:
    """
    Extracts text from PDF using pypdf.
    Raises PasswordProtectedError if the PDF is encrypted.
    """
    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        
        if reader.is_encrypted:
            raise PasswordProtectedError("The PDF document is password-protected. Please remove the password and try again.")
            
        page_count = len(reader.pages)
        text_pages = []
        
        for i, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text()
                if page_text:
                    text_pages.append(page_text)
            except Exception as e:
                if "password" in str(e).lower() or "decrypt" in str(e).lower():
                    raise PasswordProtectedError("The PDF document is password-protected. Please remove the password and try again.")
                raise DocumentParsingError(f"Error reading page {i+1}: {str(e)}")
                
        doc_text = "\n".join(text_pages)
        if not doc_text.strip():
            raise DocumentParsingError("No readable text could be extracted. The PDF might contain only scanned images.")
            
        return doc_text, page_count
        
    except PasswordProtectedError:
        raise
    except Exception as e:
        if "password" in str(e).lower() or "decrypt" in str(e).lower():
            raise PasswordProtectedError("The PDF document is password-protected. Please remove the password and try again.")
        raise DocumentParsingError(f"Failed to parse PDF document. It may be corrupted or unreadable: {str(e)}")


def extract_text_from_docx(file_path: str) -> tuple[str, int]:
    """
    Extracts text from DOCX paragraph and table elements using python-docx.
    Estimates page count based on word count.
    """
    try:
        import docx
        doc = docx.Document(file_path)
    except Exception as e:
        # Check if the error is likely due to password-protected encryption
        err_msg = str(e).lower()
        if "password" in err_msg or "encrypt" in err_msg or "compoundfile" in err_msg or "compound file" in err_msg:
            raise PasswordProtectedError("The Word document is password-protected. Please remove the password and try again.")
        raise DocumentParsingError(f"Failed to open DOCX document. It may be corrupted, encrypted, or invalid: {str(e)}")

    paragraphs = []
    # 1. Extract paragraphs
    for para in doc.paragraphs:
        p_text = para.text.strip()
        if p_text:
            paragraphs.append(p_text)

    # 2. Extract tables
    table_lines = []
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                c_text = cell.text.strip()
                if c_text:
                    cells.append(c_text)
            if cells:
                table_lines.append(" | ".join(cells))

    body_text = "\n\n".join(paragraphs)
    if table_lines:
        body_text += "\n\n--- Table Contents ---\n" + "\n".join(table_lines)

    if not body_text.strip():
        raise DocumentParsingError("The Word document contains no readable text or paragraphs.")

    # Estimate page count: standard academic page has approx 450 words
    word_count = len([w for w in body_text.split() if w.strip()])
    page_count = max(1, math.ceil(word_count / 450))

    return body_text, page_count

def find_libreoffice() -> str | None:
    """
    Attempts to locate the LibreOffice headless conversion binary (soffice).
    Checks PATH first, then standard installation directories on Windows.
    """
    # 1. Check system PATH
    soffice_path = shutil.which("soffice")
    if soffice_path:
        return soffice_path

    # 2. Check common Windows program installation folders
    if sys.platform == "win32":
        common_paths = [
            os.path.join(os.environ.get("ProgramFiles", "C:\\Program Files"), "LibreOffice", "program", "soffice.exe"),
            os.path.join(os.environ.get("ProgramFiles(x86)", "C:\\Program Files (x86)"), "LibreOffice", "program", "soffice.exe")
        ]
        for path in common_paths:
            if os.path.exists(path):
                return path

    return None

def extract_text_from_doc(file_path: str) -> tuple[str, int]:
    """
    Converts binary legacy DOC to DOCX using headless LibreOffice,
    then parses the resulting DOCX.
    """
    soffice_bin = find_libreoffice()
    if not soffice_bin:
        raise MissingDependencyError(
            "LibreOffice is not installed on this server, which is required to read legacy binary .doc files. "
            "Please install LibreOffice, or convert your document to .docx manually before uploading."
        )

    import tempfile
    
    # Create temporary directory for converted file output
    with tempfile.TemporaryDirectory() as temp_dir:
        try:
            # Execute headless soffice conversion: soffice --headless --convert-to docx --outdir <temp_dir> <file_path>
            result = subprocess.run(
                [soffice_bin, "--headless", "--convert-to", "docx", "--outdir", temp_dir, file_path],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=20  # Avoid hanging processes
            )
            
            if result.returncode != 0:
                raise DocumentParsingError(f"LibreOffice conversion failed: {result.stderr or result.stdout}")
                
        except subprocess.TimeoutExpired:
            raise DocumentParsingError("LibreOffice conversion timed out after 20 seconds. The file may be corrupt or too large.")
        except Exception as e:
            raise DocumentParsingError(f"Headless conversion execution failed: {str(e)}")

        # Find the converted .docx file in the output directory
        base_name = os.path.splitext(os.path.basename(file_path))[0]
        converted_docx = os.path.join(temp_dir, f"{base_name}.docx")

        if not os.path.exists(converted_docx):
            # Fallback search if filename was modified during conversion
            files = [f for f in os.listdir(temp_dir) if f.endswith(".docx")]
            if files:
                converted_docx = os.path.join(temp_dir, files[0])
            else:
                raise DocumentParsingError("LibreOffice completed but did not write the output docx file. It may be password-protected or corrupt.")

        # Extract text from the converted docx
        return extract_text_from_docx(converted_docx)

def extract_text(file_path: str) -> tuple[str, int]:
    """
    Main entrypoint: parses text and returns (text, page_count) from PDF, DOCX, or DOC.
    """
    if not os.path.exists(file_path):
        raise DocumentParsingError("The target document file does not exist on disk.")

    file_format = detect_file_type(file_path)

    if file_format == 'pdf':
        return extract_text_from_pdf(file_path)
    elif file_format == 'docx':
        return extract_text_from_docx(file_path)
    elif file_format == 'doc':
        return extract_text_from_doc(file_path)
    else:
        raise DocumentParsingError("Unsupported file format.")
