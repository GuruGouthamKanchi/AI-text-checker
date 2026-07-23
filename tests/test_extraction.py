import os
import sys
import docx

# Add root folder to path
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from src.document_parser import (
    extract_text,
    detect_file_type,
    DocumentParsingError,
    PasswordProtectedError,
    MissingDependencyError
)

def test_docx_extraction():
    print("Testing DOCX text and table extraction...")
    doc_path = "tests/test_temp.docx"
    doc = docx.Document()
    doc.add_heading("Deep Learning Survey", level=0)
    doc.add_paragraph("This is a paragraph about neural architectures.")
    doc.add_paragraph("Another paragraph with training parameters.")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Epochs'
    hdr_cells[1].text = 'Loss'
    row_cells = table.rows[1].cells
    row_cells[0].text = '100'
    row_cells[1].text = '0.04'
    
    doc.save(doc_path)
    
    try:
        fmt = detect_file_type(doc_path)
        assert fmt == 'docx', f"Expected docx, got {fmt}"
        
        text, page_count = extract_text(doc_path)
        assert "neural architectures" in text, "Paragraph text not found"
        assert "Epochs | Loss" in text or "100 | 0.04" in text, "Table text not found"
        assert page_count == 1, f"Expected page count 1, got {page_count}"
        print("DOCX extraction test passed!")
    finally:
        if os.path.exists(doc_path):
            os.remove(doc_path)

def test_unsupported_format():
    print("Testing unsupported file format detection...")
    txt_path = "tests/test_temp.txt"
    with open(txt_path, "w") as f:
        f.write("Hello, World!")
        
    try:
        try:
            extract_text(txt_path)
            assert False, "Expected DocumentParsingError"
        except DocumentParsingError:
            pass
        print("Unsupported format test passed!")
    finally:
        if os.path.exists(txt_path):
            os.remove(txt_path)

def test_doc_missing_libreoffice_fallback():
    print("Testing legacy .doc conversion fallback...")
    fake_doc_path = "tests/test_temp.doc"
    with open(fake_doc_path, "wb") as f:
        f.write(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 100)
        
    try:
        try:
            extract_text(fake_doc_path)
            assert False, "Expected MissingDependencyError"
        except MissingDependencyError as e:
            assert "LibreOffice is not installed" in str(e), f"Unexpected message: {str(e)}"
        print("Legacy DOC fallback test passed!")
    finally:
        if os.path.exists(fake_doc_path):
            os.remove(fake_doc_path)

if __name__ == "__main__":
    try:
        test_docx_extraction()
        test_unsupported_format()
        test_doc_missing_libreoffice_fallback()
        print("All document extraction tests passed successfully!")
    except AssertionError as e:
        print(f"Test failure: {str(e)}")
        sys.exit(1)
